"""Fixed document operations. Input is JSON data, never Python or shell code."""
import importlib.util
import json
import os
import sys
import zipfile
from pathlib import Path
from xml.sax.saxutils import escape


def read_docx(path, maximum):
    from docx import Document
    from docx.table import Table
    from docx.text.paragraph import Paragraph
    with zipfile.ZipFile(path) as archive:
        if len(archive.infolist()) > 5000 or sum(i.file_size for i in archive.infolist()) > 50_000_000:
            raise ValueError("Word 文件展开后过大")
    document = Document(path)
    blocks, used, truncated = [], 0, False
    for node in document.element.body:
        if node.tag.endswith('}p'):
            paragraph = Paragraph(node, document)
            text = paragraph.text
            if not text:
                continue
            block = {'type': 'heading' if paragraph.style.name.startswith(('Heading', 'Title')) else 'paragraph', 'text': text}
        elif node.tag.endswith('}tbl'):
            table = Table(node, document)
            block = {'type': 'table', 'rows': [[cell.text for cell in row.cells] for row in table.rows]}
        else:
            continue
        size = len(json.dumps(block, ensure_ascii=False))
        if used + size > maximum:
            truncated = True
            break
        blocks.append(block)
        used += size
    return {'format': 'docx', 'blocks': blocks, 'truncated': truncated,
            'notice': '这是段落与表格结构读取，不是 Word 的真实分页渲染。'}


def read_pdf(path, maximum, start=1, end=None):
    from pypdf import PdfReader
    reader = PdfReader(path)
    if reader.is_encrypted:
        raise ValueError('暂不读取加密 PDF，请先提供可读取的副本')
    total = len(reader.pages)
    end = min(total, end if end is not None else start + 4, start + 19)
    if start < 1 or start > total or end < start:
        raise ValueError('PDF 页码范围无效')
    pages, used, clipped = [], 0, False
    for index in range(start - 1, end):
        text = reader.pages[index].extract_text() or ''
        remaining = maximum - used
        if len(text) > remaining:
            text = text[:remaining]
            clipped = True
        pages.append({'page': index + 1, 'text': text})
        used += len(text)
        if used >= maximum:
            break
    return {'format': 'pdf', 'pageCount': total, 'pages': pages,
            'truncated': clipped or start > 1 or (pages and pages[-1]['page'] < total),
            'notice': '扫描图片中的文字需要 OCR；空文本不代表页面没有内容。'}


def validate_blocks(request):
    title = request.get('title', '文档')
    blocks = request.get('blocks', [])
    if not isinstance(title, str) or len(title) > 200 or not isinstance(blocks, list) or not 1 <= len(blocks) <= 200:
        raise ValueError('标题或内容块数量无效')
    if len(json.dumps(blocks, ensure_ascii=False)) > 100_000:
        raise ValueError('文档内容过长')
    for block in blocks:
        if not isinstance(block, dict) or block.get('type') not in ('paragraph', 'heading', 'list', 'table'):
            raise ValueError('不支持的内容块类型')
        if block['type'] == 'table':
            rows = block.get('rows', [])
            if not isinstance(rows, list) or not 1 <= len(rows) <= 100 or not isinstance(rows[0], list) or not 1 <= len(rows[0]) <= 8:
                raise ValueError('表格范围无效')
            if any(not isinstance(row, list) or len(row) != len(rows[0]) or any(not isinstance(cell, str) or len(cell) > 2000 for cell in row) for row in rows):
                raise ValueError('表格列数或内容无效')
        elif block['type'] == 'list':
            if not isinstance(block.get('items'), list) or not 1 <= len(block['items']) <= 100 or any(not isinstance(item, str) or len(item) > 4000 for item in block['items']):
                raise ValueError('列表内容无效')
        elif not isinstance(block.get('text'), str) or len(block['text']) > 12_000:
            raise ValueError('段落内容无效')
        if block['type'] == 'heading' and block.get('level', 1) not in (1, 2, 3):
            raise ValueError('标题级别无效')
    return title, blocks


def create_docx(request):
    from docx import Document
    from docx.shared import Inches, Pt, RGBColor
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
    title, blocks = validate_blocks(request)
    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin = section.bottom_margin = Inches(.8)
    section.left_margin = section.right_margin = Inches(.85)
    for name in ('Normal', 'Title', 'Heading 1', 'Heading 2', 'Heading 3'):
        style = document.styles[name]
        style.font.name = 'Arial'
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), '微软雅黑')
    document.styles['Normal'].font.size = Pt(11)
    document.styles['Normal'].paragraph_format.space_after = Pt(7)
    document.styles['Normal'].paragraph_format.line_spacing = 1.2
    document.add_heading(title, 0)
    for block in blocks:
        kind = block['type']
        if kind == 'heading':
            document.add_heading(block['text'], block.get('level', 1))
        elif kind == 'paragraph':
            document.add_paragraph(block['text'])
        elif kind == 'list':
            for item in block['items']:
                document.add_paragraph(item, style='List Bullet')
        else:
            table = document.add_table(rows=0, cols=len(block['rows'][0]))
            table.style = 'Table Grid'
            table.autofit = False
            for column in table.columns:
                column.width = Inches(6.8 / len(block['rows'][0]))
            for index, row in enumerate(block['rows']):
                word_row = table.add_row()
                if index == 0:
                    word_row._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
                cells = word_row.cells
                for cell, text in zip(cells, row):
                    cell.text = text
                    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                    borders = OxmlElement('w:tcBorders')
                    for edge in ('top', 'left', 'bottom', 'right'):
                        border = OxmlElement('w:' + edge)
                        for attr, value in [('val', 'single'), ('sz', '4'), ('color', 'D9D9D9')]:
                            border.set(qn('w:' + attr), value)
                        borders.append(border)
                    cell._tc.get_or_add_tcPr().append(borders)
                    for paragraph in cell.paragraphs:
                        for run in paragraph.runs:
                            run.font.size = Pt(10)
                            run.bold = index == 0
                    if index == 0:
                        shading = OxmlElement('w:shd'); shading.set(qn('w:fill'), 'E8EDF3')
                        cell._tc.get_or_add_tcPr().append(shading)
            document.add_paragraph()
    document.save(request['output'])
    verified = Document(request['output'])
    if not verified.paragraphs or verified.paragraphs[0].text != title:
        raise ValueError('Word 结构验证失败')
    return {'format': 'docx', 'validated': True, 'paragraphCount': len(verified.paragraphs), 'tableCount': len(verified.tables),
            'notice': '已验证 DOCX 结构；分页效果请在 Word 中查看，结构预览不代表真实 Word 排版。'}


def create_pdf(request):
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.pagesizes import letter
    from reportlab.lib import colors
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.cidfonts import UnicodeCIDFont
    from reportlab.pdfbase.ttfonts import TTFont
    title, blocks = validate_blocks(request)
    font = 'STSong-Light'
    font_path = Path(os.environ.get('WINDIR', 'C:/Windows')) / 'Fonts' / 'msyh.ttc'
    if font_path.exists():
        try:
            pdfmetrics.registerFont(TTFont('DesktopCJK', str(font_path), subfontIndex=0))
            font = 'DesktopCJK'
        except Exception:
            pass
    if font == 'STSong-Light':
        pdfmetrics.registerFont(UnicodeCIDFont(font))
    styles = getSampleStyleSheet()
    body = ParagraphStyle('ChineseBody', parent=styles['BodyText'], fontName=font, fontSize=11, leading=17, spaceAfter=8, wordWrap='CJK')
    heading = ParagraphStyle('ChineseHeading', parent=body, fontSize=15, leading=21, spaceBefore=12, spaceAfter=8, keepWithNext=True)
    title_style = ParagraphStyle('ChineseTitle', parent=heading, fontSize=21, leading=28, spaceAfter=16)
    cell_style = ParagraphStyle('ChineseCell', parent=body, fontSize=10, leading=15, spaceAfter=0)
    para = lambda text, style=body: Paragraph(escape(text).replace('\n', '<br/>'), style)
    story = [para(title, title_style)]
    for block in blocks:
        if block['type'] == 'heading':
            story.append(para(block['text'], heading))
        elif block['type'] == 'paragraph':
            story.append(para(block['text']))
        elif block['type'] == 'list':
            for item in block['items']:
                story.append(para('• ' + item))
        else:
            table = Table([[para(cell, cell_style) for cell in row] for row in block['rows']],
                          colWidths=[(letter[0] - 108) / len(block['rows'][0])] * len(block['rows'][0]), repeatRows=1)
            table.setStyle(TableStyle([('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#E8EDF3')),
                ('GRID', (0, 0), (-1, -1), .5, colors.HexColor('#D9D9D9')), ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
                ('LEFTPADDING', (0, 0), (-1, -1), 8), ('RIGHTPADDING', (0, 0), (-1, -1), 8),
                ('TOPPADDING', (0, 0), (-1, -1), 7), ('BOTTOMPADDING', (0, 0), (-1, -1), 7)]))
            story.extend([table, Spacer(1, 10)])
    def footer(canvas, doc):
        canvas.setFont(font, 9)
        canvas.drawRightString(letter[0] - 54, 30, str(doc.page))
    SimpleDocTemplate(request['output'], pagesize=letter, leftMargin=54, rightMargin=54, topMargin=48, bottomMargin=48).build(story, onFirstPage=footer, onLaterPages=footer)
    from pypdf import PdfReader
    reader = PdfReader(request['output'])
    if not reader.pages:
        raise ValueError('PDF 结构验证失败')
    return {'format': 'pdf', 'validated': True, 'pageCount': len(reader.pages)}


try:
    request = json.loads(sys.stdin.buffer.read().decode('utf-8'))
    action = request['action']
    if action == 'probe':
        result = {name: importlib.util.find_spec(name) is not None for name in ('docx', 'pypdf', 'reportlab')}
    elif action == 'read':
        maximum = max(100, min(50_000, int(request.get('maximum', 12_000))))
        if Path(request['path']).suffix.lower() == '.docx':
            result = read_docx(request['path'], maximum)
        else:
            result = read_pdf(request['path'], maximum, int(request.get('page_start', 1)), request.get('page_end'))
    elif action == 'create':
        result = create_docx(request) if request['format'] == 'docx' else create_pdf(request)
    else:
        raise ValueError('Unsupported operation')
    print(json.dumps({'ok': True, **result}, ensure_ascii=False))
except Exception as error:
    print(json.dumps({'ok': False, 'error': str(error)}, ensure_ascii=False))
    sys.exit(1)
